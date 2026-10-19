"""Turn dropped/pasted files into text the model can read."""

from __future__ import annotations

import csv
import io
import mimetypes
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Tuple

from . import config

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
SPREADSHEET_EXTS = {".xlsx", ".xlsm", ".xls"}
ARCHIVE_EXTS = {".zip"}
SKIP_IN_ARCHIVE = {"node_modules", ".git", "__pycache__", ".venv", "venv", "dist", "build", ".next"}


def kind_for(name: str, mime: str = "") -> str:
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXTS or mime.startswith("image/"):
        return "image"
    if ext == ".pdf" or mime == "application/pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in SPREADSHEET_EXTS:
        return "spreadsheet"
    if ext in ARCHIVE_EXTS:
        return "archive"
    return "text"


def _decode(data: bytes) -> Tuple[str, bool]:
    """Return (text, looks_binary)."""
    sample = data[:8192]
    if b"\x00" in sample:
        return "", True
    for enc in ("utf-8", "utf-16"):
        try:
            return data.decode(enc), False
        except UnicodeDecodeError:
            continue
    text = data.decode("latin-1")
    printable = sum(ch.isprintable() or ch in "\r\n\t" for ch in text[:8192])
    return text, printable < 0.85 * min(len(text), 8192)


def _pdf_text(path: Path) -> str:
    try:
        out = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
            capture_output=True, timeout=180, check=False,
        )
        text = out.stdout.decode("utf-8", "replace")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return "[PDF text extraction unavailable]"
    # pdftotext separates pages with form feeds; make page boundaries explicit for citations.
    pages = text.split("\f")
    if len(pages) > 1:
        text = "\n".join(f"--- page {i} ---\n{page.rstrip()}" for i, page in enumerate(pages, 1) if page.strip())
    return text.strip() or "[PDF has no extractable text layer (scanned?)]"


def _docx_text(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        return "[python-docx not installed]"
    doc = docx.Document(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        text = para.text
        if not text.strip():
            continue
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            text = "#" * min(int(level), 6) + " " + text
        parts.append(text)
    for t_index, table in enumerate(doc.tables, 1):
        parts.append(f"\n[table {t_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _spreadsheet_text(data: bytes, name: str) -> str:
    try:
        import openpyxl
    except ImportError:
        return "[openpyxl not installed]"
    if Path(name).suffix.lower() == ".xls":
        return "[legacy .xls not supported; save as .xlsx or .csv]"
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            writer.writerow(["" if v is None else v for v in row])
        out.append(f"--- sheet: {ws.title} ---\n{buf.getvalue().rstrip()}")
    return "\n\n".join(out)


def _archive_text(data: bytes) -> str:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = [i for i in zf.infolist() if not i.is_dir()]
        out.append("Archive contents:\n" + "\n".join(f"  {i.filename} ({i.file_size} bytes)" for i in names[:2000]))
        budget = config.EXTRACT_MAX_CHARS
        for info in names:
            parts = set(Path(info.filename).parts)
            if parts & SKIP_IN_ARCHIVE or info.file_size > 5 * 1024 * 1024:
                continue
            inner_kind = kind_for(info.filename)
            if inner_kind not in {"text", "pdf", "docx"}:
                continue
            raw = zf.read(info)
            if inner_kind == "text":
                text, binary = _decode(raw)
                if binary:
                    continue
            else:
                text = extract_bytes(raw, info.filename)[0]
            if len(text) > budget:
                break
            budget -= len(text)
            out.append(f"\n===== {info.filename} =====\n{text}")
    return "\n".join(out)


def extract_bytes(data: bytes, name: str, mime: str = "") -> Tuple[str, str]:
    """Return (text, kind). Images return an empty text (they go to the model as pixels)."""
    kind = kind_for(name, mime)
    try:
        if kind == "image":
            text = ""
        elif kind == "pdf":
            with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
                tmp.write(data)
                tmp.flush()
                text = _pdf_text(Path(tmp.name))
        elif kind == "docx":
            text = _docx_text(data)
        elif kind == "spreadsheet":
            text = _spreadsheet_text(data, name)
        elif kind == "archive":
            text = _archive_text(data)
        else:
            text, binary = _decode(data)
            if binary:
                kind = "binary"
                guess = mimetypes.guess_type(name)[0] or mime or "application/octet-stream"
                text = f"[binary file: {guess}, {len(data)} bytes — contents not readable as text]"
    except Exception as exc:  # never let one odd file break the upload
        text = f"[could not extract text: {exc}]"
    if len(text) > config.EXTRACT_MAX_CHARS:
        text = text[: config.EXTRACT_MAX_CHARS] + "\n[... truncated at extraction limit ...]"
    return text.replace("\r\n", "\n"), kind
